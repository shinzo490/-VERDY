# -*- coding: utf-8 -*-
"""Scripting.Dictionary 研修ブック (.xlsm) を生成する。

  python3 tools/build_xlsm.py [出力パス]

vba/*.bas を読み込んで vbaProject.bin を組み立て、
openpyxl で作ったブックにマクロを埋め込んだ .xlsm を書き出す。
"""
import datetime
import io
import os
import random
import re
import shutil
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import vbaproject

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HEAD_ROW, DATA_ROW = 7, 8

C_TITLE = "1F3864"
F_IN = PatternFill("solid", fgColor="DDEBF7")     # 入力見出し
F_OUT = PatternFill("solid", fgColor="E2EFDA")    # 出力見出し
F_MST = PatternFill("solid", fgColor="FCE4D6")    # マスタ見出し
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def head(ws, row, col, captions, fill):
    for i, c in enumerate(captions):
        cell = ws.cell(row=row, column=col + i, value=c)
        cell.font = Font(bold=True, size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX


def put_rows(ws, row, col, rows, formats=None):
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            cell = ws.cell(row=row + r, column=col + c, value=v)
            cell.border = BOX
            cell.font = Font(size=10)
            if formats and formats[c]:
                cell.number_format = formats[c]


def sheet_header(ws, title, purpose, old_way, merit, macro):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=C_TITLE)
    for row, label, text in ((2, "ねらい", purpose), (3, "従来のやり方", old_way),
                             (4, "Dictionaryの効果", merit)):
        ws.cell(row=row, column=1, value="【%s】%s" % (label, text)).font = Font(size=10)
    if macro:
        c = ws.cell(row=5, column=1, value="▶ 実行マクロ： %s　（Alt+F8 → マクロ名を選んで実行）" % macro)
        c.font = Font(bold=True, size=10, color="C00000")


def label(ws, row, col, text, color="1F3864"):
    ws.cell(row=row, column=col, value=text).font = Font(bold=True, size=10, color=color)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ----------------------------------------------------------------- データ生成
rnd = random.Random(20260827)

CUSTS = [("C001", "青葉商事"), ("C002", "北山物産"), ("C003", "さくら産業"), ("C004", "東和トレード"),
         ("C005", "みなと工業"), ("C006", "大和デンキ"), ("C007", "富士フーズ"), ("C008", "新星商会"),
         ("C009", "西口製作所"), ("C010", "湖南マテリアル")]
ITEMS = ["ボールペン", "ノート", "ファイル", "電卓", "付箋", "封筒", "テープ", "クリップ"]
BRANCHES = ["東京", "大阪", "名古屋", "福岡", "札幌", "仙台"]
PERSONS = ["山田", "佐藤", "鈴木", "高橋", "田中"]


def gen_sales(n, cols):
    rows = []
    for i in range(n):
        cust = rnd.choice(CUSTS)
        rows.append([datetime.date(2026, rnd.randint(4, 6), rnd.randint(1, 28)),
                     cust[0], cust[1], "P%03d" % rnd.randint(1, 8),
                     rnd.randrange(3, 60) * 1000])
    return [r[:cols] for r in rows]


def build_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)

    # ============================================================ 00_目次
    ws = wb.create_sheet("00_目次")
    ws.sheet_properties.tabColor = "FF1F3864"
    widths(ws, {"A": 5, "B": 20, "C": 26, "D": 46, "E": 46, "F": 24})
    ws["A1"] = "Excel VBA 中級研修　Scripting.Dictionary 実務活用ドリル"
    ws["A1"].font = Font(bold=True, size=16, color=C_TITLE)
    ws["A2"] = "対象：基本文法は書けるが、Dictionary を実務でどこに使えばよいか迷っている方"
    ws["A3"] = "使い方：各シートを開き、Alt+F8 で対応マクロを実行 → 出力エリアに結果が出ます。コードは Alt+F11、印刷用は［09_コード全文］シート。"
    ws["A5"] = "※ 開いたときに［セキュリティの警告：マクロが無効にされました］が出たら［コンテンツの有効化］をクリックしてください。"
    ws["A5"].font = Font(size=10, color="C00000", bold=True)
    ws["A4"] = "コード内の目印 … '★定義：Dictionaryの用意　／　'★書込：登録・更新　／　'★読込：取り出し・存在確認"
    for r in (2, 3, 4):
        ws.cell(row=r, column=1).font = Font(size=10)

    label(ws, 6, 1, "■ 例題一覧（Dictionary で「断然」効率的になる処理）")
    head(ws, 7, 1, ["No", "シート", "処理テーマ", "従来のやり方とその限界", "Dictionary にすると何が変わるか", "実行マクロ"], F_OUT)
    cases = [
        (1, "01_重複排除", "重複を除いた一覧を作る",
         "二重ループや COUNTIF で既出判定。1万件で最大5000万回の比較（O(n^2)）",
         "キーの重複を許さない性質で1回のループ（O(n)）。出現順も保たれる", "M01_UniqueList"),
        (2, "02_キー別集計", "キー別に件数・数量・金額を集計",
         "キー一覧を作り SUMIF/COUNTIF を項目数×キー数だけ回す（毎回全走査）",
         "明細を1回なめるだけで複数項目を同時に積み上げ。項目が増えても走査は増えない", "M02_GroupSum"),
        (3, "03_マスタ照合", "マスタから名称・単価を引く",
         "1行ごとに VLOOKUP／Find。明細×マスタ件数の照合。未登録はエラー処理が必要",
         "マスタを1回読むだけ。以降は一瞬で引ける（O(1)）。未登録は Exists で明示判定", "M03_MasterLookup"),
        (4, "04_2表突合", "前月と当月を突合し差分を出す",
         "両表を往復で総当たり検索。片方にしか無い行の拾い漏れが起きやすい",
         "両表をキー化して Exists 判定だけ。新規／削除／変更／一致を機械的に列挙できる", "M04_CompareTwoLists"),
        (5, "05_クロス集計", "支店×商品のマトリクス表を作る",
         "交点の数だけ SUMIFS（支店20×商品30＝600本）が全走査",
         "複合キー1本で1回のループ。行位置・列位置も Dictionary で持てば出力先が即決まる", "M05_CrossTab"),
        (6, "06_明細グループ化", "キー別に明細をまとめる／シート分割",
         "担当者の数だけデータを読み直す（またはフィルタを掛け直す）",
         "値に Collection を持たせ「キー→明細行番号の束」を1回で作る", "M06_GroupDetail / M06_SplitToSheets"),
        (7, "07_採番連番", "キーごとの枝番を採番する",
         "並べ替えてから前行と比較。元の並び順が崩れる。COUNTIF 方式は O(n^2)",
         "カウンタとして使えば並べ替え不要・1回のループで元の順序のまま採番", "M07_SerialNumber / M07_KeyPitfall"),
        (8, "08_速度比較", "従来方式と実測で比べる",
         "件数を2倍にすると所要時間は約4倍（O(n^2)）",
         "件数を2倍にしても約2倍（O(n)）。この“増え方の違い”が実務の差になる", "M08_Benchmark"),
    ]
    for i, row in enumerate(cases):
        for c, v in enumerate(row):
            cell = ws.cell(row=8 + i, column=1 + c, value=v)
            cell.border = BOX
            cell.font = Font(size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[8 + i].height = 34

    r = 18
    label(ws, r, 1, "■ Dictionary 早見表（すべて CreateObject(\"Scripting.Dictionary\") 前提）")
    head(ws, r + 1, 1, ["", "記述", "意味・注意点", "", "", ""], F_OUT)
    ref = [
        ("定義", 'Set dic = CreateObject("Scripting.Dictionary")', "参照設定が不要な遅延バインディング。研修・配布用はこちらが安全"),
        ("定義", "dic.CompareMode = 1", "1=vbTextCompare。大文字小文字を区別しない。※キー登録前に設定すること"),
        ("書込", "dic.Add key, value", "新規追加。既存キーを Add すると実行時エラー（重複チェックに使える）"),
        ("書込", "dic(key) = value", "無ければ追加、有れば上書き。集計はこの書き方が簡潔"),
        ("読込", "dic.Exists(key)", "存在確認。True/False。キーは増えない"),
        ("読込", "v = dic(key)", "値の取り出し。※未登録キーを読むと空の値で自動追加される（落とし穴）"),
        ("読込", "dic.Keys / dic.Items", "登録順に並んだ配列を返す。For Each で回す"),
        ("読込", "dic.Count", "登録件数。ユニーク件数そのもの"),
        ("削除", "dic.Remove key / dic.RemoveAll", "1件削除／全削除"),
    ]
    for i, (kind, code, note) in enumerate(ref):
        ws.cell(row=r + 2 + i, column=1, value=kind).font = Font(size=10, bold=True)
        ws.cell(row=r + 2 + i, column=2, value=code).font = Font(size=10, name="Consolas")
        ws.cell(row=r + 2 + i, column=3, value=note).font = Font(size=10)
        ws.cell(row=r + 2 + i, column=3).alignment = Alignment(wrap_text=True)
        for c in range(1, 4):
            ws.cell(row=r + 2 + i, column=c).border = BOX

    r = 30
    label(ws, r, 1, "■ 中級者がつまずく5点")
    tips = [
        "キーの型を揃える … \"001\" と 1 は別のキー。CStr() で文字列に統一するのが定石",
        "値が配列のときは直接書き換えられない … v = dic(k) → 更新 → dic(k) = v と書き戻す（例2）",
        "値がオブジェクト（Collection など）なら直接 .Add できる … 参照を持っているため（例6）",
        "未登録キーを読むとキーが増える … 存在確認は必ず Exists（例7の M07_KeyPitfall で実演）",
        "複合キーの区切り文字はデータに現れない文字を使う … vbTab が安全（例5）",
    ]
    for i, t in enumerate(tips):
        ws.cell(row=r + 1 + i, column=1, value="(%d)" % (i + 1)).font = Font(size=10, bold=True)
        ws.cell(row=r + 1 + i, column=2, value=t).font = Font(size=10)

    r = 38
    label(ws, r, 1, "■ 速さの目安（1件あたりの処理を1とした比較イメージ）")
    head(ws, r + 1, 1, ["", "件数", "COUNTIF / VLOOKUP 方式", "Dictionary 方式", "", ""], F_OUT)
    for i, n in enumerate((1000, 5000, 10000, 50000)):
        ws.cell(row=r + 2 + i, column=2, value=n).border = BOX
        ws.cell(row=r + 2 + i, column=3, value="%s 回の比較" % format(n * n // 2, ",")).border = BOX
        ws.cell(row=r + 2 + i, column=4, value="%s 回の登録" % format(n, ",")).border = BOX
        for c in (2, 3, 4):
            ws.cell(row=r + 2 + i, column=c).font = Font(size=10)

    ws.freeze_panes = "A7"

    # ============================================================ 01_重複排除
    ws = wb.create_sheet("01_重複排除")
    widths(ws, {"A": 12, "B": 14, "C": 16, "D": 12, "E": 12, "F": 3, "G": 6, "H": 14, "I": 16})
    sheet_header(ws, "例1　重複の除去（ユニーク一覧の作成）",
                 "売上明細から得意先の一覧を、重複なし・出現順で作る。",
                 "二重ループや COUNTIF で既出かどうかを判定する（件数の2乗に比例して遅くなる）。",
                 "キーの重複を許さない性質をそのまま使い、1回のループで完了。登録順も保たれる。",
                 "M01_UniqueList")
    label(ws, 6, 1, "■ 入力データ（売上明細）")
    label(ws, 6, 7, "■ 出力（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["受注日", "得意先コード", "得意先名", "商品コード", "金額"], F_IN)
    head(ws, HEAD_ROW, 7, ["No", "得意先コード", "得意先名"], F_OUT)
    put_rows(ws, DATA_ROW, 1, gen_sales(40, 5), ["yyyy/mm/dd", None, None, None, "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 02_キー別集計
    ws = wb.create_sheet("02_キー別集計")
    widths(ws, {"A": 10, "B": 14, "C": 8, "D": 12, "E": 3, "F": 10, "G": 8, "H": 10, "I": 12})
    sheet_header(ws, "例2　キー別の集計（件数・数量計・金額計を一度に）",
                 "支店ごとに件数・数量計・金額計をまとめて集計する。",
                 "支店一覧を作り、SUMIF と COUNTIF を支店数×集計項目数だけ回す（毎回データ全体を走査）。",
                 "明細を1回なめるだけで3項目を同時に積み上げる。集計項目が増えても走査回数は変わらない。",
                 "M02_GroupSum")
    label(ws, 6, 1, "■ 入力データ（売上明細）")
    label(ws, 6, 6, "■ 出力（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["支店", "商品", "数量", "金額"], F_IN)
    head(ws, HEAD_ROW, 6, ["支店", "件数", "数量計", "金額計"], F_OUT)
    rows = []
    for _ in range(60):
        qty = rnd.randint(1, 30)
        rows.append([rnd.choice(BRANCHES), rnd.choice(ITEMS), qty, qty * rnd.choice([120, 250, 480, 900])])
    put_rows(ws, DATA_ROW, 1, rows, [None, None, "#,##0", "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 03_マスタ照合
    ws = wb.create_sheet("03_マスタ照合")
    widths(ws, {"A": 12, "B": 8, "C": 3, "D": 12, "E": 16, "F": 8, "G": 3,
                "H": 12, "I": 8, "J": 16, "K": 8, "L": 12})
    sheet_header(ws, "例3　マスタ照合（VLOOKUP／Find の繰り返しを置き換える）",
                 "明細の商品コードから商品名・単価を引き、金額を計算する。未登録コードは明示する。",
                 "1行ごとに VLOOKUP や Find でマスタを検索（明細件数×マスタ件数）。未登録はエラー処理が必要。",
                 "マスタを最初に1回だけ読み込めば、以降は一瞬で引ける。未登録判定も Exists で正面から書ける。",
                 "M03_MasterLookup")
    label(ws, 6, 1, "■ 入力（受注明細）")
    label(ws, 6, 4, "■ 商品マスタ", "833C00")
    label(ws, 6, 8, "■ 出力（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["商品コード", "数量"], F_IN)
    head(ws, HEAD_ROW, 4, ["商品コード", "商品名", "単価"], F_MST)
    head(ws, HEAD_ROW, 8, ["商品コード", "数量", "商品名", "単価", "金額"], F_OUT)
    grades = ["エコ", "スリム", "プロ", "標準"]
    master = [["P%03d" % i,
               "%s %s" % (rnd.choice(grades), ITEMS[(i - 1) % len(ITEMS)]),
               rnd.choice([100, 180, 250, 380, 500, 780])] for i in range(1, 21)]
    put_rows(ws, DATA_ROW, 4, master, [None, None, "#,##0"])
    codes = ["P%03d" % i for i in range(1, 21)] + ["P098", "P099"]      # 末尾2つはマスタに無い
    detail = [[rnd.choice(codes), rnd.randint(1, 20)] for _ in range(50)]
    detail[7][0], detail[31][0] = "P098", "P099"                        # 未登録を必ず含める
    put_rows(ws, DATA_ROW, 1, detail, [None, "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 04_2表突合
    ws = wb.create_sheet("04_2表突合")
    widths(ws, {"A": 12, "B": 10, "C": 3, "D": 12, "E": 10, "F": 3,
                "G": 12, "H": 10, "I": 10, "J": 10, "K": 16})
    sheet_header(ws, "例4　2つの表の突合（新規／削除／変更／一致）",
                 "前月リストと当月リストを突合し、増えた・減った・変わった・同じ を区分する。",
                 "前月→当月、当月→前月と往復で総当たり検索。片方にしか無い行を拾い漏らしやすい。",
                 "両表をキー化すれば Exists の判定だけで済み、拾い漏れが構造的に起きない。",
                 "M04_CompareTwoLists")
    label(ws, 6, 1, "■ 前月リスト")
    label(ws, 6, 4, "■ 当月リスト")
    label(ws, 6, 7, "■ 出力（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["商品コード", "数量"], F_IN)
    head(ws, HEAD_ROW, 4, ["商品コード", "数量"], F_IN)
    head(ws, HEAD_ROW, 7, ["商品コード", "前月数量", "当月数量", "差分", "区分"], F_OUT)
    prev_codes = ["A%03d" % i for i in range(1, 26)]
    prev = [[c, rnd.randint(5, 80)] for c in prev_codes]
    curr = []
    for c, q in prev:
        if c in ("A004", "A011", "A019"):          # 当月で消えたコード
            continue
        curr.append([c, q if rnd.random() < 0.5 else q + rnd.choice([-6, -3, 4, 9])])
    curr += [["A101", 12], ["A102", 40]]           # 当月の新規
    put_rows(ws, DATA_ROW, 1, prev, [None, "#,##0"])
    put_rows(ws, DATA_ROW, 4, curr, [None, "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 05_クロス集計
    ws = wb.create_sheet("05_クロス集計")
    widths(ws, {"A": 10, "B": 14, "C": 12, "D": 3, "E": 12,
                "F": 11, "G": 11, "H": 11, "I": 11, "J": 11, "K": 11, "L": 11})
    sheet_header(ws, "例5　複合キーによるクロス集計（支店 × 商品）",
                 "明細から「行＝支店・列＝商品」のマトリクス表を作る。",
                 "支店一覧・商品一覧を作り、交点の数だけ SUMIFS を置く（交点が増えるほど全走査が増える）。",
                 "複合キー（支店＋区切り＋商品）1本で1回のループ。行位置・列位置も Dictionary で管理する。",
                 "M05_CrossTab")
    label(ws, 6, 1, "■ 入力データ（売上明細）")
    label(ws, 6, 5, "■ 出力（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["支店", "商品", "金額"], F_IN)
    rows = [[rnd.choice(BRANCHES[:5]), rnd.choice(ITEMS[:6]), rnd.randrange(5, 90) * 1000] for _ in range(60)]
    put_rows(ws, DATA_ROW, 1, rows, [None, None, "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 06_明細グループ化
    ws = wb.create_sheet("06_明細グループ化")
    widths(ws, {"A": 10, "B": 16, "C": 14, "D": 12, "E": 3, "F": 24, "G": 16, "H": 14, "I": 12})
    sheet_header(ws, "例6　1つのキーに複数行をぶら下げる（グループ化・シート分割）",
                 "明細を担当者ごとにまとめて出力する。さらに担当者別シートへ分割する。",
                 "担当者の数だけデータ全体をループし直す（またはオートフィルタを掛け直してコピー）。",
                 "値に Collection を入れて「キー→明細行番号の束」を1回のループで作る。読むのは1回だけ。",
                 "M06_GroupDetail　※分割は M06_SplitToSheets／後始末は M06_DeleteSplitSheets")
    label(ws, 6, 1, "■ 入力データ（売上明細）")
    label(ws, 6, 6, "■ 出力（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["担当者", "得意先", "商品", "金額"], F_IN)
    rows = [[rnd.choice(PERSONS), rnd.choice(CUSTS)[1], rnd.choice(ITEMS), rnd.randrange(3, 50) * 1000]
            for _ in range(40)]
    put_rows(ws, DATA_ROW, 1, rows, [None, None, None, "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 07_採番連番
    ws = wb.create_sheet("07_採番連番")
    widths(ws, {"A": 12, "B": 14, "C": 12, "D": 3, "E": 14, "F": 8, "G": 14, "H": 12,
                "I": 3, "J": 8, "K": 34, "L": 30})
    sheet_header(ws, "例7　キーごとの連番（枝番）採番と、Dictionary の落とし穴",
                 "受注明細に得意先ごとの通し番号を振り、\"得意先コード-001\" の管理番号を作る。",
                 "得意先コードで並べ替えてから前行と比較（元の並び順が崩れる）。COUNTIF 方式は O(n^2)。",
                 "カウンタとして使えば並べ替え不要。1回のループで元の順序を保ったまま採番できる。",
                 "M07_SerialNumber　※落とし穴の実演は M07_KeyPitfall")
    label(ws, 6, 1, "■ 入力データ（受注明細）")
    label(ws, 6, 5, "■ 出力（マクロが書き込みます）", "C00000")
    label(ws, 6, 10, "■ 落とし穴の実演（M07_KeyPitfall）", "833C00")
    head(ws, HEAD_ROW, 1, ["受注日", "得意先コード", "金額"], F_IN)
    head(ws, HEAD_ROW, 5, ["得意先コード", "枝番", "管理番号", "金額"], F_OUT)
    rows = []
    for _ in range(40):
        cust = rnd.choice(CUSTS[:8])
        rows.append([datetime.date(2026, rnd.randint(4, 6), rnd.randint(1, 28)), cust[0],
                     rnd.randrange(3, 60) * 1000])
    put_rows(ws, DATA_ROW, 1, rows, ["yyyy/mm/dd", None, "#,##0"])
    ws.freeze_panes = "A8"

    # ============================================================ 08_速度比較
    ws = wb.create_sheet("08_速度比較")
    widths(ws, {"A": 14, "B": 14, "C": 12, "D": 44, "E": 22, "F": 14, "G": 12, "H": 30})
    ws["A1"] = "例8　速度比較（従来方式 vs Dictionary）"
    ws["A1"].font = Font(bold=True, size=14, color=C_TITLE)
    ws["A2"] = "【ねらい】ダミーデータを自動生成し、ユニーク件数を数える処理を COUNTIF 方式と Dictionary 方式で実測する。"
    ws["A3"] = "【見どころ】件数を2倍にすると COUNTIF 方式は約4倍（O(n^2)）、Dictionary 方式は約2倍（O(n)）。"
    for r in (2, 3):
        ws.cell(row=r, column=1).font = Font(size=10)
    ws["B4"] = "テスト件数 →"
    ws["B4"].font = Font(bold=True, size=10)
    ws["B4"].alignment = Alignment(horizontal="right")
    ws["C4"] = 5000
    ws["C4"].font = Font(bold=True, size=11, color="C00000")
    ws["C4"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["C4"].border = BOX
    ws["D4"] = "← まず 5000 で実行。30000 を超えると従来方式は数分かかります。"
    ws["D4"].font = Font(size=10)
    ws["A5"] = "▶ 実行マクロ： M08_Benchmark　（Alt+F8 → マクロ名を選んで実行）"
    ws["A5"].font = Font(bold=True, size=10, color="C00000")
    label(ws, 6, 1, "■ 生成データ（マクロが書き込みます）")
    label(ws, 6, 5, "■ 計測結果（マクロが書き込みます）", "C00000")
    head(ws, HEAD_ROW, 1, ["商品コード"], F_IN)
    head(ws, HEAD_ROW, 5, ["方式", "ユニーク件数", "所要秒数", "備考"], F_OUT)
    ws.freeze_panes = "A8"

    # ============================================================ 09_コード全文
    ws = wb.create_sheet("09_コード全文")
    ws.sheet_properties.tabColor = "FF7F7F7F"
    widths(ws, {"A": 118})
    ws["A1"] = "コード全文（印刷・回覧用）　※ 実物は Alt+F11 → 標準モジュール で確認・編集できます"
    ws["A1"].font = Font(bold=True, size=12, color=C_TITLE)
    r = 3
    mono = Font(name="MS Gothic", size=9)
    for fname in sorted(os.listdir(os.path.join(ROOT, "vba"))):
        if not fname.lower().endswith(".bas"):
            continue
        with io.open(os.path.join(ROOT, "vba", fname), encoding="utf-8") as f:
            lines = f.read().replace("\r\n", "\n").split("\n")
        title = lines[0].split('"')[1]
        c = ws.cell(row=r, column=1, value="●―――― %s ――――" % title)
        c.font = Font(bold=True, size=11, color="C00000")
        r += 1
        for line in lines[1:]:
            cell = ws.cell(row=r, column=1, value=line if line else None)
            cell.font = mono
            cell.alignment = Alignment(vertical="center")
            r += 1
        r += 2
    ws.freeze_panes = "A3"

    for i, sh in enumerate(wb.worksheets):
        sh.sheet_properties.codeName = "Sheet%d" % (i + 1)
        sh.sheet_view.showGridLines = True
    wb.code_name = "ThisWorkbook"
    wb.save(path)
    return len(wb.worksheets)


# ------------------------------------------------------------- xlsm 化
def make_xlsm(xlsx_path, xlsm_path, vba_bin):
    zin = zipfile.ZipFile(xlsx_path)
    with zipfile.ZipFile(xlsm_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                t = data.decode("utf-8")
                t = t.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                              "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
                t = t.replace("</Types>",
                              '<Override PartName="/xl/vbaProject.bin" '
                              'ContentType="application/vnd.ms-office.vbaProject"/></Types>')
                data = t.encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                t = data.decode("utf-8")
                used = re.findall(r'Id="rId(\d+)"', t)
                new_id = "rId%d" % (max(int(x) for x in used) + 1 if used else 1)
                t = t.replace("</Relationships>",
                              '<Relationship Id="%s" Type="http://schemas.microsoft.com/office/2006/'
                              'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>' % new_id)
                data = t.encode("utf-8")
            elif item.filename == "xl/workbook.xml":
                t = data.decode("utf-8")
                if "<workbookPr" in t:
                    if "codeName=" not in t.split("<workbookPr", 1)[1].split(">", 1)[0]:
                        t = t.replace("<workbookPr", '<workbookPr codeName="ThisWorkbook"', 1)
                else:
                    t = t.replace("<sheets>", '<workbookPr codeName="ThisWorkbook"/><sheets>', 1)
                data = t.encode("utf-8")
            zout.writestr(item, data)
        zout.write(vba_bin, "xl/vbaProject.bin")
    zin.close()


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        ROOT, "Dictionary実務活用ドリル.xlsm")
    tmp_dir = os.path.join(ROOT, ".build")
    os.makedirs(tmp_dir, exist_ok=True)
    xlsx = os.path.join(tmp_dir, "book.xlsx")
    vba_bin = os.path.join(tmp_dir, "vbaProject.bin")

    n_sheets = build_workbook(xlsx)

    modules = [vbaproject.doc_module("ThisWorkbook")]
    modules += [vbaproject.doc_module("Sheet%d" % i) for i in range(1, n_sheets + 1)]
    for name in sorted(os.listdir(os.path.join(ROOT, "vba"))):
        if name.lower().endswith(".bas"):
            with open(os.path.join(ROOT, "vba", name), encoding="utf-8") as f:
                modules.append(vbaproject.std_module_from_bas(f.read()))
    vbaproject.build(modules, vba_bin)

    make_xlsm(xlsx, out, vba_bin)
    shutil.rmtree(tmp_dir, ignore_errors=True)
    print("wrote %s (%.1f KB, %d modules)" % (out, os.path.getsize(out) / 1024.0, len(modules)))


if __name__ == "__main__":
    main()
